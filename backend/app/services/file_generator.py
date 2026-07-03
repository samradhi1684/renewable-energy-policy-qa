import json
import os
import uuid

from docx import Document
from openpyxl import Workbook

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
)

GENERATED_DIR = "generated"

os.makedirs(
    GENERATED_DIR,
    exist_ok=True,
)


def generate_file(
    answer: str,
    file_type: str,
):

    file_id = str(uuid.uuid4())

    if file_type == "pdf":

        filename = f"{file_id}.pdf"

        path = os.path.join(
            GENERATED_DIR,
            filename,
        )

        doc = SimpleDocTemplate(path)

        doc.build(
            [
                Paragraph(answer)
            ]
        )

        return filename

    if file_type == "docx":

        filename = f"{file_id}.docx"

        path = os.path.join(
            GENERATED_DIR,
            filename,
        )

        doc = Document()

        doc.add_paragraph(
            answer
        )

        doc.save(path)

        return filename

    if file_type == "json":

        filename = f"{file_id}.json"

        path = os.path.join(
            GENERATED_DIR,
            filename,
        )

        with open(
            path,
            "w",
            encoding="utf-8",
        ) as f:

            json.dump(
                {
                    "answer": answer
                },
                f,
                indent=2,
            )

        return filename

    if file_type == "md":

        filename = f"{file_id}.md"

        path = os.path.join(
            GENERATED_DIR,
            filename,
        )

        with open(
            path,
            "w",
            encoding="utf-8",
        ) as f:

            f.write(answer)

        return filename

    if file_type == "xlsx":

        filename = f"{file_id}.xlsx"

        path = os.path.join(
            GENERATED_DIR,
            filename,
        )

        wb = Workbook()

        ws = wb.active

        ws.append(
            ["Answer"]
        )

        ws.append(
            [answer]
        )

        wb.save(path)

        return filename

    return None